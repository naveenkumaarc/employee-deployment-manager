from __future__ import annotations

from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Any

from PySide6.QtCore import QDate, Qt, Signal
from PySide6.QtGui import QColor
from PySide6.QtWidgets import (
    QFileDialog,
    QFrame,
    QHBoxLayout,
    QLabel,
    QMessageBox,
    QTabWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from .database import Database
from .dialogs import DayEntriesDialog, DeploymentDialog
from .theme import COLORS
from .widgets import (
    AdaptiveComboBox,
    MetricCard,
    MonthCalendar,
    PageHeading,
    SearchBox,
    fill_combo,
    format_currency,
    make_button,
    make_date_edit,
    make_table,
    set_combo_value,
)

DETAIL_COLUMNS = [
    "Date",
    "Employee",
    "Customer",
    "Project",
    "Site",
    "Work type",
    "Time",
    "Status",
    "Day basis",
    "Amount / day",
    "Payment",
    "Notes",
]


def _month_bounds(day: QDate | None = None) -> tuple[QDate, QDate]:
    value = day or QDate.currentDate()
    first = QDate(value.year(), value.month(), 1)
    return first, first.addMonths(1).addDays(-1)


def _iso(value: QDate) -> str:
    return value.toString("yyyy-MM-dd")


def _payment_filter_values(mode: str | None) -> tuple[bool | None, bool | None]:
    if mode == "day_basis":
        return True, None
    if mode == "pending":
        return True, False
    if mode == "completed":
        return True, True
    return None, None


def _add_payment_filters(combo: AdaptiveComboBox) -> None:
    combo.addItem("All deployments", "all")
    combo.addItem("Day-basis only", "day_basis")
    combo.addItem("Payment pending", "pending")
    combo.addItem("Payment completed", "completed")


def _populate_detail_table(table, rows: list[dict[str, Any]]) -> None:
    table.setRowCount(len(rows))
    status_colors = {
        "Scheduled": QColor("#536177"),
        "Confirmed": QColor(COLORS["primary"]),
        "Completed": QColor(COLORS["success"]),
        "Cancelled": QColor(COLORS["danger"]),
    }
    for row_index, row in enumerate(rows):
        values = [
            QDate.fromString(row["work_date"], "yyyy-MM-dd").toString("dd MMM yyyy"),
            row["employee_name"],
            row["customer_name"],
            (
                f"{row['project_code']} — {row['project_name']}"
                if row.get("project_code")
                else row["project_name"]
            ),
            row["site_name"]
            + (f", {row['site_city']}" if row.get("site_city") else ""),
            row["work_type"],
            f"{row['start_time']}–{row['end_time']}",
            row["status"],
            "Yes" if row.get("day_basis") else "No",
            format_currency(row.get("day_rate")) if row.get("day_basis") else "—",
            (
                "Completed" if row.get("payment_completed") else "Pending"
            )
            if row.get("day_basis")
            else "—",
            row.get("notes") or "",
        ]
        for column, value in enumerate(values):
            item = QTableWidgetItem(str(value))
            if column == 0:
                item.setData(Qt.ItemDataRole.UserRole, row["id"])
            if column == 7:
                item.setForeground(
                    status_colors.get(row["status"], QColor(COLORS["text"]))
                )
                font = item.font()
                font.setBold(True)
                item.setFont(font)
            if column == 9:
                item.setTextAlignment(
                    Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter
                )
            if column == 10 and row.get("day_basis"):
                completed = bool(row.get("payment_completed"))
                item.setForeground(
                    QColor(COLORS["success"] if completed else COLORS["amber"])
                )
                font = item.font()
                font.setBold(True)
                item.setFont(font)
            table.setItem(row_index, column, item)
    table.resizeRowsToContents()


class CalendarPage(QWidget):
    dataChanged = Signal()
    requestDirectory = Signal()

    def __init__(self, database: Database, parent: QWidget | None = None):
        super().__init__(parent)
        self.database = database

        layout = QVBoxLayout(self)
        layout.setContentsMargins(26, 22, 26, 22)
        layout.setSpacing(14)

        title_row = QHBoxLayout()
        title_row.addWidget(
            PageHeading(
                "Deployment calendar",
                "Click an empty date to add; click an occupied date to view or edit its entries.",
            ),
            1,
        )
        self.new_button = make_button("＋  New deployment", role="Primary")
        self.new_button.clicked.connect(self.add_selected)
        title_row.addWidget(self.new_button)
        layout.addLayout(title_row)

        toolbar = QFrame()
        toolbar.setObjectName("ToolbarCard")
        tool_layout = QHBoxLayout(toolbar)
        tool_layout.setContentsMargins(12, 9, 12, 9)
        tool_layout.setSpacing(8)

        previous = make_button("‹", tooltip="Previous month")
        previous.setFixedWidth(42)
        today = make_button("Today")
        following = make_button("›", tooltip="Next month")
        following.setFixedWidth(42)
        previous.clicked.connect(lambda: self.calendar.showPreviousMonth())
        today.clicked.connect(self._show_today)
        following.clicked.connect(lambda: self.calendar.showNextMonth())
        tool_layout.addWidget(previous)
        tool_layout.addWidget(today)
        tool_layout.addWidget(following)

        self.month_label = QLabel()
        self.month_label.setObjectName("SectionTitle")
        self.month_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        tool_layout.addStretch()
        tool_layout.addWidget(self.month_label)
        tool_layout.addStretch()

        directory = make_button("Employees && directory")
        directory.clicked.connect(self.requestDirectory.emit)
        tool_layout.addWidget(directory)
        layout.addWidget(toolbar)

        self.calendar = MonthCalendar()
        self.calendar.clicked.connect(self._date_clicked)
        self.calendar.pageChanged.connect(self._page_changed)
        layout.addWidget(self.calendar, 1)

        self.refresh()

    def _show_today(self) -> None:
        self.calendar.setSelectedDate(QDate.currentDate())
        self.calendar.showToday()
        self.refresh()

    def _page_changed(self, _year: int, _month: int) -> None:
        self.refresh()

    def refresh(self) -> None:
        year, month = self.calendar.yearShown(), self.calendar.monthShown()
        self.month_label.setText(QDate(year, month, 1).toString("MMMM yyyy"))
        self.calendar.set_events(self.database.month_events(year, month))

    def add_selected(self) -> None:
        if not self.database.required_data_ready():
            QMessageBox.information(
                self,
                "Set up the directory first",
                "Add at least one active employee, customer, project and site before "
                "creating a deployment.",
            )
            self.requestDirectory.emit()
            return
        if DeploymentDialog(
            self.database,
            initial_date=self.calendar.selectedDate(),
            parent=self,
        ).exec():
            self.refresh()
            self.dataChanged.emit()

    def _date_clicked(self, day: QDate) -> None:
        self.calendar.setSelectedDate(day)
        if self.calendar.events_for(day):
            dialog = DayEntriesDialog(self.database, day, self)
            dialog.exec()
            if dialog.changed:
                self.refresh()
                self.dataChanged.emit()
        else:
            self.add_selected()


class HistoryPage(QWidget):
    dataChanged = Signal()

    def __init__(self, database: Database, parent: QWidget | None = None):
        super().__init__(parent)
        self.database = database
        self.rows: list[dict[str, Any]] = []

        layout = QVBoxLayout(self)
        layout.setContentsMargins(26, 22, 26, 22)
        layout.setSpacing(14)
        layout.addWidget(
            PageHeading(
                "Deployment history",
                "Search, filter, edit or remove individual deployment records.",
            )
        )

        filters = QFrame()
        filters.setObjectName("Card")
        filter_layout = QVBoxLayout(filters)
        filter_layout.setContentsMargins(14, 13, 14, 13)
        filter_layout.setSpacing(10)

        first, last = _month_bounds()
        row_one = QHBoxLayout()
        self.date_from = make_date_edit(first)
        self.date_to = make_date_edit(last)
        self.employee = AdaptiveComboBox()
        self.customer = AdaptiveComboBox()
        self.status = AdaptiveComboBox()
        self.status.addItem("All statuses", None)
        for value in ("Scheduled", "Confirmed", "Completed", "Cancelled"):
            self.status.addItem(value, value)
        row_one.addWidget(QLabel("From"))
        row_one.addWidget(self.date_from)
        row_one.addWidget(QLabel("To"))
        row_one.addWidget(self.date_to)
        row_one.addWidget(self.employee, 1)
        row_one.addWidget(self.customer, 1)
        row_one.addWidget(self.status, 1)
        filter_layout.addLayout(row_one)

        row_two = QHBoxLayout()
        self.project = AdaptiveComboBox()
        self.site = AdaptiveComboBox()
        self.payment_filter = AdaptiveComboBox()
        _add_payment_filters(self.payment_filter)
        self.search = SearchBox("Search employee, company, project, site or notes")
        apply_button = make_button("Apply filters", role="Primary")
        clear_button = make_button("Reset")
        apply_button.clicked.connect(self.refresh)
        clear_button.clicked.connect(self.reset_filters)
        self.search.returnPressed.connect(self.refresh)
        row_two.addWidget(self.project, 1)
        row_two.addWidget(self.site, 1)
        row_two.addWidget(self.payment_filter, 1)
        row_two.addWidget(self.search, 2)
        row_two.addWidget(clear_button)
        row_two.addWidget(apply_button)
        filter_layout.addLayout(row_two)
        layout.addWidget(filters)

        controls = QHBoxLayout()
        self.count_label = QLabel()
        self.count_label.setObjectName("Muted")
        controls.addWidget(self.count_label)
        controls.addStretch()
        copy = make_button("Copy selected")
        edit = make_button("Edit selected")
        delete = make_button("Delete selected", role="Danger")
        copy.clicked.connect(self.copy_selected)
        edit.clicked.connect(self.edit_selected)
        delete.clicked.connect(self.delete_selected)
        controls.addWidget(copy)
        controls.addWidget(edit)
        controls.addWidget(delete)
        layout.addLayout(controls)

        self.table = make_table(DETAIL_COLUMNS)
        self.table.doubleClicked.connect(lambda _index: self.edit_selected())
        layout.addWidget(self.table, 1)

        for widget in (
            self.date_from,
            self.date_to,
            self.employee,
            self.customer,
            self.project,
            self.site,
            self.status,
            self.payment_filter,
        ):
            if hasattr(widget, "dateChanged"):
                widget.dateChanged.connect(lambda *_: self.refresh())
            else:
                widget.currentIndexChanged.connect(lambda *_: self.refresh())
        self.refresh_options()
        self.refresh()

    def refresh_options(self) -> None:
        fill_combo(
            self.employee,
            self.database.list_master("employee"),
            include_all=True,
            all_label="All employees",
        )
        fill_combo(
            self.customer,
            self.database.list_master("customer"),
            include_all=True,
            all_label="All customers",
        )
        fill_combo(
            self.project,
            self.database.list_master("project"),
            include_all=True,
            all_label="All projects",
        )
        fill_combo(
            self.site,
            self.database.list_master("site"),
            include_all=True,
            all_label="All sites",
        )

    def reset_filters(self) -> None:
        first, last = _month_bounds()
        self.date_from.setDate(first)
        self.date_to.setDate(last)
        for combo in (
            self.employee,
            self.customer,
            self.project,
            self.site,
            self.status,
            self.payment_filter,
        ):
            combo.setCurrentIndex(0)
        self.search.clear()
        self.refresh()

    def _query(self) -> list[dict[str, Any]]:
        day_basis, payment_completed = _payment_filter_values(
            self.payment_filter.currentData()
        )
        return self.database.list_deployments(
            date_from=_iso(self.date_from.date()),
            date_to=_iso(self.date_to.date()),
            employee_id=self.employee.currentData(),
            customer_id=self.customer.currentData(),
            project_id=self.project.currentData(),
            site_id=self.site.currentData(),
            status=self.status.currentData(),
            day_basis=day_basis,
            payment_completed=payment_completed,
            search=self.search.text(),
        )

    def refresh(self) -> None:
        if self.date_from.date() > self.date_to.date():
            return
        self.rows = self._query()
        _populate_detail_table(self.table, self.rows)
        total_amount = sum(
            float(row.get("day_rate") or 0)
            for row in self.rows
            if row.get("day_basis")
        )
        pending_amount = sum(
            float(row.get("day_rate") or 0)
            for row in self.rows
            if row.get("day_basis") and not row.get("payment_completed")
        )
        self.count_label.setText(
            f"{len(self.rows)} deployment record{'s' if len(self.rows) != 1 else ''}"
            f"  •  Day-basis total {format_currency(total_amount)}"
            f"  •  Pending {format_currency(pending_amount)}"
        )

    def selected_id(self) -> int | None:
        selected = self.table.selectionModel().selectedRows()
        if not selected:
            return None
        item = self.table.item(selected[0].row(), 0)
        return int(item.data(Qt.ItemDataRole.UserRole)) if item else None

    def edit_selected(self) -> None:
        deployment_id = self.selected_id()
        if deployment_id is None:
            QMessageBox.information(self, "Select a deployment", "Select a row to edit.")
            return
        if DeploymentDialog(
            self.database,
            deployment_id=deployment_id,
            parent=self,
        ).exec():
            self.refresh()
            self.dataChanged.emit()

    def copy_selected(self) -> None:
        deployment_id = self.selected_id()
        if deployment_id is None:
            QMessageBox.information(
                self, "Select a deployment", "Select a row to copy."
            )
            return
        if DeploymentDialog(
            self.database,
            copy_from_id=deployment_id,
            parent=self,
        ).exec():
            self.refresh()
            self.dataChanged.emit()

    def delete_selected(self) -> None:
        deployment_id = self.selected_id()
        if deployment_id is None:
            QMessageBox.information(self, "Select a deployment", "Select a row first.")
            return
        answer = QMessageBox.question(
            self,
            "Delete deployment",
            "Delete the selected deployment permanently?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if answer == QMessageBox.StandardButton.Yes:
            self.database.delete_deployment(deployment_id)
            self.refresh()
            self.dataChanged.emit()


class ReportsPage(QWidget):
    def __init__(self, database: Database, parent: QWidget | None = None):
        super().__init__(parent)
        self.database = database
        self.rows: list[dict[str, Any]] = []
        self.summary_rows: list[dict[str, Any]] = []

        layout = QVBoxLayout(self)
        layout.setContentsMargins(26, 22, 26, 22)
        layout.setSpacing(14)
        layout.addWidget(
            PageHeading(
                "Deployment reports",
                "Review day-basis payments by employee, month, date, project or customer.",
            )
        )

        filters = QFrame()
        filters.setObjectName("Card")
        filter_layout = QVBoxLayout(filters)
        filter_layout.setContentsMargins(14, 13, 14, 13)
        filter_layout.setSpacing(10)

        current = QDate.currentDate()
        row_one = QHBoxLayout()
        self.group_by = AdaptiveComboBox()
        self.group_by.addItem("Employee-wise", "employee")
        self.group_by.addItem("Month-wise", "month")
        self.group_by.addItem("Date-wise", "date")
        self.group_by.addItem("Project-wise", "project")
        self.group_by.addItem("Customer-wise", "customer")
        self.date_from = make_date_edit(QDate(current.year(), 1, 1))
        self.date_to = make_date_edit(current)
        self.employee = AdaptiveComboBox()
        self.customer = AdaptiveComboBox()
        row_one.addWidget(QLabel("Report"))
        row_one.addWidget(self.group_by)
        row_one.addWidget(QLabel("From"))
        row_one.addWidget(self.date_from)
        row_one.addWidget(QLabel("To"))
        row_one.addWidget(self.date_to)
        row_one.addWidget(self.employee, 1)
        row_one.addWidget(self.customer, 1)
        filter_layout.addLayout(row_one)

        row_two = QHBoxLayout()
        self.project = AdaptiveComboBox()
        self.site = AdaptiveComboBox()
        self.status = AdaptiveComboBox()
        self.status.addItem("All statuses", None)
        for value in ("Scheduled", "Confirmed", "Completed", "Cancelled"):
            self.status.addItem(value, value)
        self.payment_filter = AdaptiveComboBox()
        _add_payment_filters(self.payment_filter)
        generate = make_button("Generate report", role="Primary")
        export = make_button("Export Excel")
        generate.clicked.connect(self.generate)
        export.clicked.connect(self.export_excel)
        row_two.addWidget(self.project, 1)
        row_two.addWidget(self.site, 1)
        row_two.addWidget(self.status, 1)
        row_two.addWidget(self.payment_filter, 1)
        row_two.addStretch()
        row_two.addWidget(export)
        row_two.addWidget(generate)
        filter_layout.addLayout(row_two)
        layout.addWidget(filters)

        metrics = QHBoxLayout()
        self.total_metric = MetricCard("Deployments", accent=COLORS["primary"])
        self.payment_metric = MetricCard("Day-basis total", accent=COLORS["teal"])
        self.completed_metric = MetricCard("Payment completed", accent=COLORS["success"])
        self.pending_metric = MetricCard("Payment pending", accent=COLORS["amber"])
        for card in (
            self.total_metric,
            self.payment_metric,
            self.completed_metric,
            self.pending_metric,
        ):
            metrics.addWidget(card, 1)
        layout.addLayout(metrics)

        tabs = QTabWidget()
        self.summary_table = make_table(
            [
                "Group",
                "Deployments",
                "Work dates",
                "Employees",
                "Total amount",
                "Completed",
                "Pending",
            ]
        )
        self.details_table = make_table(DETAIL_COLUMNS)
        tabs.addTab(self.summary_table, "Summary")
        tabs.addTab(self.details_table, "Detailed records")
        layout.addWidget(tabs, 1)

        self.refresh_options()
        self.generate()

    def refresh_options(self) -> None:
        fill_combo(
            self.employee,
            self.database.list_master("employee"),
            include_all=True,
            all_label="All employees",
        )
        fill_combo(
            self.customer,
            self.database.list_master("customer"),
            include_all=True,
            all_label="All customers",
        )
        fill_combo(
            self.project,
            self.database.list_master("project"),
            include_all=True,
            all_label="All projects",
        )
        fill_combo(
            self.site,
            self.database.list_master("site"),
            include_all=True,
            all_label="All sites",
        )

    def _query(self) -> list[dict[str, Any]]:
        day_basis, payment_completed = _payment_filter_values(
            self.payment_filter.currentData()
        )
        return self.database.list_deployments(
            date_from=_iso(self.date_from.date()),
            date_to=_iso(self.date_to.date()),
            employee_id=self.employee.currentData(),
            customer_id=self.customer.currentData(),
            project_id=self.project.currentData(),
            site_id=self.site.currentData(),
            status=self.status.currentData(),
            day_basis=day_basis,
            payment_completed=payment_completed,
        )

    def _group_label(self, row: dict[str, Any], mode: str) -> str:
        if mode == "employee":
            return row["employee_name"]
        if mode == "project":
            code = f"{row['project_code']} — " if row.get("project_code") else ""
            return code + row["project_name"]
        if mode == "customer":
            return row["customer_name"]
        if mode == "date":
            return QDate.fromString(row["work_date"], "yyyy-MM-dd").toString(
                "dd MMM yyyy"
            )
        value = datetime.strptime(row["work_date"][:7], "%Y-%m")
        return value.strftime("%B %Y")

    def _summarize(self, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        mode = str(self.group_by.currentData())
        groups: OrderedDict[str, dict[str, Any]] = OrderedDict()
        for row in rows:
            label = self._group_label(row, mode)
            bucket = groups.setdefault(
                label,
                {
                    "group": label,
                    "deployments": 0,
                    "dates": set(),
                    "employees": set(),
                    "total_amount": 0.0,
                    "completed_amount": 0.0,
                    "pending_amount": 0.0,
                },
            )
            bucket["deployments"] += 1
            bucket["dates"].add(row["work_date"])
            bucket["employees"].add(row["employee_id"])
            if row.get("day_basis"):
                amount = float(row.get("day_rate") or 0)
                bucket["total_amount"] += amount
                if row.get("payment_completed"):
                    bucket["completed_amount"] += amount
                else:
                    bucket["pending_amount"] += amount

        result: list[dict[str, Any]] = []
        for bucket in groups.values():
            result.append(
                {
                    "group": bucket["group"],
                    "deployments": bucket["deployments"],
                    "dates": len(bucket["dates"]),
                    "employees": len(bucket["employees"]),
                    "total_amount": round(bucket["total_amount"], 2),
                    "completed_amount": round(bucket["completed_amount"], 2),
                    "pending_amount": round(bucket["pending_amount"], 2),
                }
            )
        if mode not in {"month", "date"}:
            result.sort(key=lambda item: item["group"].casefold())
        return result

    def generate(self) -> None:
        if self.date_from.date() > self.date_to.date():
            QMessageBox.warning(
                self, "Invalid date range", "The From date must be before the To date."
            )
            return
        self.rows = self._query()
        self.summary_rows = self._summarize(self.rows)

        self.summary_table.setRowCount(len(self.summary_rows))
        keys = (
            "group",
            "deployments",
            "dates",
            "employees",
            "total_amount",
            "completed_amount",
            "pending_amount",
        )
        for row_index, row in enumerate(self.summary_rows):
            for column, key in enumerate(keys):
                value = format_currency(row[key]) if column >= 4 else str(row[key])
                item = QTableWidgetItem(value)
                if 0 < column < 4:
                    item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                elif column >= 4:
                    item.setTextAlignment(
                        Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter
                    )
                if column == 5:
                    item.setForeground(QColor(COLORS["success"]))
                elif column == 6:
                    item.setForeground(QColor(COLORS["amber"]))
                self.summary_table.setItem(row_index, column, item)
        self.summary_table.resizeRowsToContents()
        _populate_detail_table(self.details_table, self.rows)

        self.total_metric.set_value(len(self.rows))
        total_amount = sum(
            float(row.get("day_rate") or 0)
            for row in self.rows
            if row.get("day_basis")
        )
        completed_amount = sum(
            float(row.get("day_rate") or 0)
            for row in self.rows
            if row.get("day_basis") and row.get("payment_completed")
        )
        pending_amount = total_amount - completed_amount
        self.payment_metric.set_value(format_currency(total_amount))
        self.completed_metric.set_value(format_currency(completed_amount))
        self.pending_metric.set_value(format_currency(pending_amount))

    def _filter_description(self) -> list[tuple[str, str]]:
        return [
            ("Report type", self.group_by.currentText()),
            (
                "Period",
                f"{self.date_from.date().toString('dd MMM yyyy')} to "
                f"{self.date_to.date().toString('dd MMM yyyy')}",
            ),
            ("Employee", self.employee.currentText()),
            ("Customer", self.customer.currentText()),
            ("Project", self.project.currentText()),
            ("Site", self.site.currentText()),
            ("Status", self.status.currentText()),
            ("Payment", self.payment_filter.currentText()),
        ]

    def export_excel(self) -> None:
        self.generate()
        if not self.rows:
            QMessageBox.information(self, "Nothing to export", "The report has no records.")
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            QMessageBox.critical(
                self,
                "Excel support unavailable",
                "Install the required packages using:\n\npip install -r requirements.txt",
            )
            return

        suggested = (
            Path.home()
            / "Documents"
            / f"Deployment_Report_{datetime.now():%Y%m%d_%H%M}.xlsx"
        )
        path, _ = QFileDialog.getSaveFileName(
            self, "Export deployment report", str(suggested), "Excel workbook (*.xlsx)"
        )
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"

        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"
        summary.sheet_view.showGridLines = False
        summary.merge_cells("A1:G1")
        summary["A1"] = "Deployment Desk — Deployment & Payment Report"
        summary["A1"].font = Font(size=18, bold=True, color="172033")
        summary["A1"].alignment = Alignment(vertical="center")
        summary.row_dimensions[1].height = 34

        row_cursor = 3
        for label, value in self._filter_description():
            summary.cell(row_cursor, 1, label).font = Font(bold=True, color="536177")
            summary.cell(row_cursor, 2, value)
            row_cursor += 1
        row_cursor += 1

        headers = [
            "Group",
            "Deployments",
            "Work dates",
            "Employees",
            "Total amount",
            "Payment completed",
            "Payment pending",
        ]
        header_fill = PatternFill("solid", fgColor="526DFE")
        header_font = Font(bold=True, color="FFFFFF")
        thin = Side(style="thin", color="E3E9F2")
        for column, header in enumerate(headers, 1):
            cell = summary.cell(row_cursor, column, header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        data_start = row_cursor + 1
        for record in self.summary_rows:
            values = [
                record["group"],
                record["deployments"],
                record["dates"],
                record["employees"],
                record["total_amount"],
                record["completed_amount"],
                record["pending_amount"],
            ]
            for column, value in enumerate(values, 1):
                cell = summary.cell(row_cursor + 1, column, value)
                cell.border = Border(bottom=thin)
                if column > 1:
                    cell.alignment = Alignment(horizontal="center")
                if column >= 5:
                    cell.number_format = '₹#,##0.00'
            row_cursor += 1
        totals = [
            "GRAND TOTAL",
            len(self.rows),
            len({row["work_date"] for row in self.rows}),
            len({row["employee_id"] for row in self.rows}),
            sum(
                float(row.get("day_rate") or 0)
                for row in self.rows
                if row.get("day_basis")
            ),
            sum(
                float(row.get("day_rate") or 0)
                for row in self.rows
                if row.get("day_basis") and row.get("payment_completed")
            ),
            sum(
                float(row.get("day_rate") or 0)
                for row in self.rows
                if row.get("day_basis") and not row.get("payment_completed")
            ),
        ]
        row_cursor += 1
        for column, value in enumerate(totals, 1):
            cell = summary.cell(row_cursor, column, value)
            cell.font = Font(bold=True, color="172033")
            cell.fill = PatternFill("solid", fgColor="EEF1FF")
            cell.border = Border(top=thin, bottom=thin)
            if column >= 5:
                cell.number_format = '₹#,##0.00'
        summary.freeze_panes = f"A{data_start}"
        summary.auto_filter.ref = f"A{data_start - 1}:G{row_cursor - 1}"
        for index, width in enumerate((34, 15, 15, 13, 18, 20, 18), 1):
            summary.column_dimensions[get_column_letter(index)].width = width

        detail = workbook.create_sheet("Deployments")
        detail.sheet_view.showGridLines = False
        detail_headers = [
            "Date",
            "Employee code",
            "Employee",
            "Designation",
            "Customer",
            "Project code",
            "Project",
            "Site",
            "City",
            "Start",
            "End",
            "Work type",
            "Status",
            "Day basis",
            "Amount per employee / day",
            "Payment status",
            "Notes",
        ]
        for column, header in enumerate(detail_headers, 1):
            cell = detail.cell(1, column, header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for row_index, record in enumerate(self.rows, 2):
            values = [
                record["work_date"],
                record.get("employee_code") or "",
                record["employee_name"],
                record.get("designation") or "",
                record["customer_name"],
                record.get("project_code") or "",
                record["project_name"],
                record["site_name"],
                record.get("site_city") or "",
                record["start_time"],
                record["end_time"],
                record["work_type"],
                record["status"],
                "Yes" if record.get("day_basis") else "No",
                float(record.get("day_rate") or 0) if record.get("day_basis") else 0,
                (
                    "Completed" if record.get("payment_completed") else "Pending"
                )
                if record.get("day_basis")
                else "Not applicable",
                record.get("notes") or "",
            ]
            for column, value in enumerate(values, 1):
                cell = detail.cell(row_index, column, value)
                cell.border = Border(bottom=thin)
                cell.alignment = Alignment(vertical="top", wrap_text=column == 17)
                if column == 15:
                    cell.number_format = '₹#,##0.00'
        detail.freeze_panes = "A2"
        detail.auto_filter.ref = detail.dimensions
        widths = (
            13,
            15,
            23,
            22,
            25,
            15,
            28,
            25,
            18,
            10,
            10,
            16,
            14,
            12,
            25,
            18,
            42,
        )
        for index, width in enumerate(widths, 1):
            detail.column_dimensions[get_column_letter(index)].width = width

        try:
            workbook.save(path)
        except OSError as exc:
            QMessageBox.critical(self, "Export failed", str(exc))
            return
        QMessageBox.information(
            self,
            "Report exported",
            f"The Excel report was saved successfully:\n\n{path}",
        )
